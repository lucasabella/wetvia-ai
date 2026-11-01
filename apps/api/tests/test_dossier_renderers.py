"""Real-artifact tests for the packaged concept dossier renderers."""

from __future__ import annotations

import asyncio
import json
from datetime import date
from decimal import Decimal
from importlib.resources import files
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

from wetvia.dossier import CONCEPT_DISCLOSURE, DakkapelDossierService
from wetvia.dossier.renderers import render_docx, render_markdown, render_pdf
from wetvia.llm.adapters.fake import FakeProvider
from wetvia.rules import DakkapelInvoer, Dakvlak, Gebiedsbescherming, Monumentrelatie

_SNAPSHOT = Path(__file__).parent / "fixtures" / "dossier" / "dakkapel.md"
_CUSTOM_PROPERTIES_NS = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
_VARIANTS_NS = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"


def _facts(*, measurements: bool = True, unknown_context: bool = False) -> DakkapelInvoer:
    return DakkapelInvoer(
        dakvlak=Dakvlak.ACHTER,
        dakvlak_naar_openbaar_toegankelijk_gebied=None if unknown_context else False,
        plat_dak=None if unknown_context else True,
        hoogte_m=Decimal("1.75") if measurements else None,
        onderzijde_boven_dakvoet_m=Decimal("0.75") if measurements else None,
        bovenzijde_onder_daknok_m=Decimal("0.75") if measurements else None,
        linker_zijkant_tot_dakvlakrand_m=Decimal("0.75") if measurements else None,
        rechter_zijkant_tot_dakvlakrand_m=Decimal("0.75") if measurements else None,
        monumentrelatie=(
            Monumentrelatie.ONBEKEND if unknown_context else Monumentrelatie.GEEN_MONUMENT
        ),
        gebiedsbescherming=(
            Gebiedsbescherming.ONBEKEND if unknown_context else Gebiedsbescherming.GEEN
        ),
    )


def _dossier(
    *,
    location: str = "Voorbeeldstraat 4, Utrecht",
    description: str = "Plaatsen van een dakkapel aan de achterzijde.",
    measurements: bool = True,
    unknown_context: bool = False,
):
    async def create():
        provider = FakeProvider([json.dumps({"projectomschrijving": description})])
        service = DakkapelDossierService(
            provider,
            today=lambda: date(2026, 11, 1),
        )
        return await service.create(
            _facts(measurements=measurements, unknown_context=unknown_context),
            location=location,
        )

    return asyncio.run(create())


def test_markdown_matches_golden_snapshot() -> None:
    dossier = _dossier()
    expected = _SNAPSHOT.read_text(encoding="utf-8")

    assert render_markdown(dossier) == expected


def test_markdown_escapes_hostile_html_in_location_and_model_text() -> None:
    hostile = '<img src="https://example.invalid/track" onerror="alert(1)">'
    output = render_markdown(_dossier(location=hostile, description="<script>alert('x')</script>"))

    assert "&lt;img" in output
    assert "&lt;script" in output
    assert "<img" not in output
    assert "<script" not in output


def test_packaged_templates_are_available_as_package_resources() -> None:
    template_package = files("wetvia.dossier.templates")

    assert template_package.joinpath("dakkapel.docx").is_file()
    assert template_package.joinpath("dakkapel.html").is_file()


def test_docx_is_a_real_watermarked_artifact_with_ooxml_ai_properties() -> None:
    dossier = _dossier()
    content = render_docx(dossier)

    document = Document(BytesIO(content))
    assert document.core_properties.author == "Wetvia-AI"
    assert document.core_properties.last_modified_by == "Wetvia-AI"
    assert CONCEPT_DISCLOSURE in document.sections[0].footer.paragraphs[0].text

    with ZipFile(BytesIO(content)) as archive:
        header = archive.read("word/header1.xml").decode("utf-8")
        document_xml = archive.read("word/document.xml").decode("utf-8")
        custom_xml = ET.fromstring(archive.read("docProps/custom.xml"))
        content_types = archive.read("[Content_Types].xml").decode("utf-8")
        relationships = archive.read("_rels/.rels").decode("utf-8")

    assert "rotation:315" in header
    assert 'string="CONCEPT"' in header
    assert "Projectomschrijving" in document_xml or "Plaatsen van een dakkapel" in document_xml
    assert "/docProps/custom.xml" in content_types
    assert "custom-properties" in relationships
    properties = {
        element.attrib["name"]: element
        for element in custom_xml.findall(f"{{{_CUSTOM_PROPERTIES_NS}}}property")
    }
    assert properties["ai_generated"].find(f"{{{_VARIANTS_NS}}}bool").text == "true"
    assert properties["generator"].find(f"{{{_VARIANTS_NS}}}lpwstr").text == "Wetvia-AI"


def test_docx_escapes_hostile_location_and_keeps_unknown_measurements_honest() -> None:
    hostile = '<img src="https://example.invalid/track" onerror="alert(1)">'
    content = render_docx(
        _dossier(
            location=hostile,
            description="<script>alert('x')</script>",
            measurements=False,
            unknown_context=True,
        )
    )
    document = Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert hostile in text
    assert "<script>alert('x')</script>" in text
    assert "onbekend" in text
    assert "Plat dak: onbekend" in text
    assert "Naar openbaar toegankelijk gebied: onbekend" in text
    assert "Hoogte: onbekend" in text
    assert "Onderzijde boven dakvoet: onbekend" in text
    assert "Bovenzijde onder daknok: onbekend" in text
    assert "Linkerzijkant tot dakvlakrand: onbekend" in text
    assert "Rechterzijkant tot dakvlakrand: onbekend" in text
    assert "None m" not in text

    with ZipFile(BytesIO(content)) as archive:
        relationships = archive.read("word/_rels/document.xml.rels").decode("utf-8")
    assert "example.invalid" not in relationships


def test_pdf_is_a_real_watermarked_artifact_with_producer_and_xmp_ai_marking() -> None:
    content = render_pdf(_dossier(measurements=False, unknown_context=True))
    reader = PdfReader(BytesIO(content))
    metadata = reader.metadata
    assert metadata is not None
    assert metadata.get("/Producer") == "Wetvia-AI"
    assert metadata.get("/ai_generated") == "true"
    assert reader.xmp_metadata is not None
    xmp = reader.xmp_metadata.stream.get_data()
    assert b'wetvia:ai_generated="true"' in xmp
    assert b'wetvia:generator="Wetvia-AI"' in xmp
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "CONCEPT" in text
    assert "onbekend" in text
    assert "None m" not in text
    assert "Plat dak: onbekend" in text


def test_pdf_autoescapes_user_content_and_does_not_fetch_resources(monkeypatch) -> None:
    from wetvia.dossier.renderers import pdf as pdf_renderer

    hostile = '<img src="https://example.invalid/track" onerror="alert(1)">'
    fetches: list[str] = []

    def record_forbidden_fetch(url: str, *_args: Any, **_kwargs: Any) -> Any:
        fetches.append(url)
        raise AssertionError(f"resource fetch was attempted: {url}")

    monkeypatch.setattr(pdf_renderer, "_deny_resource_fetch", record_forbidden_fetch)
    content = render_pdf(
        _dossier(
            location=hostile,
            description="<script>alert('x')</script>",
            unknown_context=True,
        )
    )
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)

    assert hostile in text
    assert "<script>alert('x')</script>" in text
    assert fetches == []
